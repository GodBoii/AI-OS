"""Exercise the real toolkit, Chromium renderer, and editable PPTX output."""

import copy
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ppt_tools import TEMPLATES, PresentationTools
from presentation_spec import content_issues, parse_outline


@pytest.fixture
def toolkit():
    return PresentationTools(user_id=None, session_id=None, message_id=None)


def example_slides(image_path: str) -> list[dict]:
    return [
        {
            "type": "title",
            "title": "A practical path to faster customer support",
            "subtitle": "Operations proposal | Internal review",
            "kicker": "Customer operations",
        },
        {
            "type": "content",
            "title": "Start with the queue that customers notice",
            "bullets": [
                "Route urgent requests to a named owner",
                "Track repeat questions before adding automation",
                "Keep a human review for unusual cases",
            ],
            "callout": "Fix the handoff before adding more software.",
        },
        {
            "type": "two_column",
            "title": "A shared queue makes ownership visible",
            "left_title": "Today",
            "right_title": "Pilot",
            "left_content": [
                "Requests spread across inboxes",
                "Priority depends on who sees the message",
                "Handoffs lack a clear owner",
            ],
            "right_content": [
                "One queue for incoming work",
                "A published priority rule",
                "A named owner for each handoff",
            ],
        },
        {
            "type": "chart",
            "title": "The sample queue falls after triage",
            "chart": {
                "type": "bar",
                "data": [
                    {"label": "Week 1", "value": 84},
                    {"label": "Week 2", "value": 68},
                    {"label": "Week 3", "value": 49},
                    {"label": "Week 4", "value": 31},
                ],
            },
            "source": "Illustrative test data, not a business forecast",
            "callout": "Review the oldest requests first.",
        },
        {
            "type": "metrics",
            "title": "Measure service quality as well as speed",
            "metrics": [
                {"value": "24h", "label": "Illustrative response target"},
                {"value": "1", "label": "Owner per request"},
                {"value": "Weekly", "label": "Review cadence"},
            ],
        },
        {
            "type": "diagram",
            "title": "Run the pilot through four review gates",
            "steps": [
                {
                    "title": "Define",
                    "detail": "Choose the queue and publish the priority rule.",
                },
                {"title": "Route", "detail": "Assign owners and log each handoff."},
                {"title": "Review", "detail": "Read a sample of completed requests."},
                {"title": "Decide", "detail": "Expand only if service quality holds."},
            ],
        },
        {
            "type": "table",
            "title": "Keep responsibility close to the work",
            "table": [
                ["Role", "Owns", "Review"],
                ["Operations", "Queue and routing", "Daily"],
                ["Team lead", "Quality and exceptions", "Weekly"],
                ["Sponsor", "Scope and next decision", "End of pilot"],
            ],
        },
        {
            "type": "image",
            "title": "Make the working queue visible",
            "image_path": image_path,
            "caption": "Synthetic queue illustration for toolkit verification",
            "bullets": [
                "Show owner, priority and age",
                "Keep exceptions visible",
                "Review the same view as a team",
            ],
        },
        {
            "type": "section",
            "title": "Protect service quality as volume grows",
            "subtitle": "The next decision depends on evidence from the pilot.",
        },
        {
            "type": "chart",
            "title": "Track the direction of weekly backlog",
            "chart": {
                "type": "line",
                "data": [
                    {"label": "W1", "value": 84},
                    {"label": "W2", "value": 68},
                    {"label": "W3", "value": 75},
                    {"label": "W4", "value": 31},
                ],
            },
            "source": "Illustrative test data",
        },
        {
            "type": "chart",
            "title": "Compare the workload across teams",
            "chart": {
                "type": "column",
                "data": [
                    {"label": "Team A", "value": 21},
                    {"label": "Team B", "value": 34},
                    {"label": "Team C", "value": 28},
                ],
            },
            "source": "Illustrative test data",
        },
        {
            "type": "content",
            "title": "Define exceptions before expanding",
            "bullets": [
                "Escalate requests with sensitive account details",
                "Keep an audit trail for changed priorities",
                "Let the owner ask for a second review",
            ],
        },
        {
            "type": "two_column",
            "title": "Scale only when the pilot earns it",
            "left_title": "Expand when",
            "right_title": "Pause when",
            "left_content": [
                "Owners can resolve most handoffs",
                "Quality reviews show consistent answers",
            ],
            "right_content": [
                "Urgent cases still wait in the queue",
                "Automation adds more review work",
            ],
        },
        {
            "type": "metrics",
            "title": "Keep the next review focused",
            "metrics": [
                {"value": "4", "label": "Weekly review gates"},
                {"value": "3", "label": "Measures to compare"},
                {"value": "1", "label": "Decision at the end"},
            ],
            "callout": "Targets and counts are examples for this test deck.",
        },
        {
            "type": "closing",
            "title": "Approve a bounded pilot and review the evidence",
            "bullets": [
                "Choose one queue and a responsible team",
                "Publish the review criteria before launch",
                "Return with findings and a clear decision",
            ],
            "callout": "Decision requested: appoint the pilot owner.",
        },
    ]


def to_outline(slides: list[dict]) -> str:
    sections = []
    for slide in slides:
        lines = [f"## [{slide['type']}] {slide['title']}"]
        for key, value in slide.items():
            if key in {"type", "title"}:
                continue
            if key == "bullets":
                lines.extend(f"- {item}" for item in value)
            else:
                lines.append(
                    f"{key}: {json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value}"
                )
        sections.append("\n".join(lines))
    return "\n\n".join(sections)


@pytest.fixture
def slides(tmp_path):
    image = tmp_path / "queue.png"
    Image.new("RGB", (900, 500), "#1B5299").save(image)
    return example_slides(str(image))


def test_outline_roundtrip(slides):
    assert parse_outline(to_outline(slides)) == slides
    assert content_issues(slides) == []


@pytest.mark.parametrize("template", list(TEMPLATES))
def test_every_layout_in_every_template(toolkit, slides, template):
    result = toolkit.validate_presentation("Customer support pilot", slides, template)
    assert result["ok"], result
    assert result["layout_issues"] == [], result


@pytest.mark.parametrize(
    "count,template", [(5, "executive"), (10, "startup_pitch"), (15, "aetheria_modern")]
)
def test_real_export_and_editable_content(toolkit, slides, count, template):
    selected = slides[: count - 1] + [slides[-1]]
    selected[0]["notes"] = "Explain why this pilot starts with one queue."
    result = toolkit.create_presentation_from_outline(
        "Customer support pilot",
        to_outline(selected),
        template,
        expected_slide_count=count,
    )
    assert result["ok"], result
    assert result["quality"]["status"] == "verified", result
    assert result["quality"]["issues"] == []
    assert "data:image" not in json.dumps(result)
    assert (
        len(json.dumps(result)) < 6000
    )  # No image bytes or repeated geometry in model context.
    path = Path(result["data"]["local_path"])
    prs = Presentation(path)
    assert len(prs.slides) == count
    assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.001
    assert "one queue" in prs.slides[0].notes_slide.notes_text_frame.text
    for slide, original in zip(prs.slides, selected):
        text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert original["title"] in text
        for bullet in original.get("bullets", []):
            assert bullet in text
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 100
            assert shape.top + shape.height <= prs.slide_height + 100
    manifest = Path(result["quality"]["preview_manifest"])
    assert len(list((manifest.parent / "previews").glob("*.jpg"))) == count
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith(".xml"):
                ElementTree.fromstring(archive.read(name))


@pytest.mark.parametrize(
    "change",
    [
        {"title": "A" * 101},
        {"bullets": ["x"] * 5},
        {"type": "invented"},
        {
            "type": "chart",
            "chart": {
                "data": [
                    {"label": "A", "value": float("nan")},
                    {"label": "B", "value": 1},
                ]
            },
        },
        {"type": "table", "table": [["A", "B"], ["missing column"]]},
        {"type": "image"},
        {"type": "metrics", "metrics": [{"value": "1", "label": "x"}] * 5},
    ],
)
def test_invalid_content_is_blocked(toolkit, change):
    slide = {"type": "content", "title": "A claim", "bullets": ["A point"], **change}
    result = toolkit.create_presentation("Invalid", [slide])
    assert not result["ok"]
    assert result["issues"]


def test_draft_batch_is_atomic_and_repairable(toolkit, slides):
    draft_id = toolkit.start_presentation_draft("Pilot", expected_slide_count=5)[
        "data"
    ]["draft_id"]
    bad = copy.deepcopy(slides[:2])
    bad[-1]["bullets"] *= 3
    assert not toolkit.add_presentation_slides(draft_id, bad)["ok"]
    assert toolkit._presentation_drafts[draft_id]["slides"] == []
    assert toolkit.add_presentation_slides(draft_id, slides[:5])["ok"]
    assert not toolkit.add_presentation_slides(draft_id, slides[:1])["ok"]
    assert toolkit.replace_presentation_slide(
        draft_id,
        2,
        {"type": "content", "title": "Revised claim", "bullets": ["Revised evidence"]},
    )["ok"]
    assert not toolkit.replace_presentation_slide(draft_id, 0, slides[1])["ok"]
    result = toolkit.finalize_presentation_draft(draft_id)
    assert result["ok"], result
    assert draft_id not in toolkit._presentation_drafts


def test_brief_never_invents_data(toolkit):
    result = toolkit.create_presentation_from_brief(
        "Support", "Make a deck about support", 5
    )
    assert not result["ok"]


def test_exact_count_and_missing_image(toolkit, slides):
    assert not toolkit.create_presentation_from_outline(
        "Pilot", to_outline(slides[:5]), expected_slide_count=10
    )["ok"]
    assert not toolkit.create_presentation(
        "Missing",
        [{"type": "image", "title": "Missing image", "image_path": "not-a-file.png"}],
    )["ok"]


def test_known_bad_geometry_cannot_export(toolkit):
    html = (
        '<div class="slide-container" style="width:1920px;height:1080px"><div data-object="true" data-object-type="textbox" style="position:absolute;left:100px;top:100px;width:100px;height:10px;font-size:36px">'
        + "Too much text " * 100
        + "</div></div>"
    )
    result = toolkit.create_presentation(
        "Bad contract", [{"type": "html", "title": "Bad", "html": html}]
    )
    assert not result["ok"], result
    assert result["issues"]


def test_thumbnails_use_event_only(toolkit, slides):
    class Socket:
        def __init__(self):
            self.events = []

        def emit(self, name, payload, **kwargs):
            self.events.append((name, payload))

    socket = Socket()
    toolkit.socketio = socket
    toolkit.session_id = "test-session"
    result = toolkit.create_presentation("Pilot", slides[:1])
    assert result["ok"], result
    assert "data:image" not in json.dumps(result)
    assert "data:image" in json.dumps(socket.events)
    assert (
        result["metadata"]["output_id"] == socket.events[0][1]["metadata"]["output_id"]
    )


@pytest.mark.parametrize("count", [0, -1, 31, True, "10"])
def test_draft_count_is_never_silently_clamped(toolkit, count):
    assert not toolkit.start_presentation_draft("Pilot", expected_slide_count=count)[
        "ok"
    ]


def test_sources_survive_non_chart_layouts(toolkit):
    slides = [
        {
            "type": "content",
            "title": "Use the supplied findings",
            "bullets": ["A concrete observation"],
            "source": "Internal study, September 2026",
        }
    ]
    result = toolkit.create_presentation("Study", slides)
    assert result["ok"], result
    prs = Presentation(result["data"]["local_path"])
    assert any(
        slides[0]["source"] in shape.text
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
    )


@pytest.mark.parametrize("chart_type", ["bar", "column", "line"])
def test_negative_and_zero_chart_values(toolkit, chart_type):
    result = toolkit.create_presentation(
        "Variance",
        [
            {
                "type": "chart",
                "title": "Variance crosses zero",
                "chart": {
                    "type": chart_type,
                    "data": [
                        {"label": "A", "value": -20},
                        {"label": "B", "value": 0},
                        {"label": "C", "value": 15},
                    ],
                },
                "source": "Illustrative test data",
            }
        ],
    )
    assert result["ok"], result
    assert result["quality"]["status"] == "verified"


def test_broken_rendered_image_is_blocked(toolkit, tmp_path):
    invalid_image = tmp_path / "broken.png"
    invalid_image.write_bytes(b"not an image")
    result = toolkit.create_presentation(
        "Image",
        [{"type": "image", "title": "Broken asset", "image_path": str(invalid_image)}],
    )
    assert not result["ok"]
    assert any(issue["type"] == "broken_image" for issue in result["issues"])


def test_layout_names_are_validated_as_text():
    assert content_issues([{"type": {}, "title": "Invalid"}])[0]["field"] == "type"


def test_evidence_alias_keeps_the_chart(toolkit):
    slide = {
        "type": "evidence",
        "title": "Compare the supplied values",
        "chart": {
            "type": "bar",
            "data": [{"label": "A", "value": 5}, {"label": "B", "value": 10}],
        },
        "source": "Illustrative test data",
    }
    result = toolkit.create_presentation("Evidence", [slide])
    assert result["ok"], result
    prs = Presentation(result["data"]["local_path"])
    text = "\n".join(
        shape.text for shape in prs.slides[0].shapes if shape.has_text_frame
    )
    assert "A" in text and "10" in text


def test_cover_metrics_keep_readable_labels(toolkit):
    result = toolkit.create_presentation(
        "Pilot",
        [
            {
                "type": "title",
                "title": "A bounded pilot",
                "subtitle": "Review the evidence before expanding",
                "metrics": [
                    {"value": "4", "label": "Review gates"},
                    {"value": "1", "label": "Owner"},
                ],
            }
        ],
    )
    assert result["ok"], result
    assert result["quality"]["issues"] == []


def test_invisible_contract_text_is_blocked(toolkit):
    html = '<div class="slide-container"><div data-object="true" data-object-type="textbox" style="position:absolute;left:100px;top:100px;width:600px;height:100px;color:#FFFFFF;font-size:36px">Invisible text</div></div>'
    result = toolkit.create_presentation(
        "Contrast", [{"type": "html", "title": "Invisible", "html": html}]
    )
    assert not result["ok"]
    assert any(issue["type"] == "low_contrast" for issue in result["issues"])


def test_dense_content_is_refused_without_truncation(toolkit):
    point = "A" * 150
    result = toolkit.create_presentation(
        "Dense",
        [
            {
                "type": "content",
                "title": "Dense evidence",
                "bullets": [point] * 4,
                "callout": "A narrow text area needs readable evidence.",
            }
        ],
    )
    assert not result["ok"], result
    assert result["issues"]
