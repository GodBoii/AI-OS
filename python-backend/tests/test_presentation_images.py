import io
import json
import sys
from pathlib import Path
from unittest.mock import Mock

import pytest
from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import presentation_images
from openrouter_image_client import GeneratedImage, ImageGenerationError
from ppt_tools import PresentationTools, build_presentation_agent
from presentation_images import PresentationImageTools, prepare_slide_images


@pytest.fixture
def provider(monkeypatch):
    buffer = io.BytesIO()
    Image.new("RGB", (96, 64), "#1B5299").save(buffer, format="PNG")
    call = Mock(
        return_value=GeneratedImage(
            buffer.getvalue(), "image/png", "test-image-model", 0.01
        )
    )
    monkeypatch.setattr(presentation_images, "generate_openrouter_image", call)
    return call


def test_generation_returns_real_local_png_and_reuses_asset(provider):
    tool = PresentationImageTools(api_key="test-key")
    first = tool.create_presentation_image("An operations desk", "executive")
    data = json.loads(first.content)
    assert data["ok"] and first.images
    asset = data["data"]
    assert asset["width"] == 96 and asset["height"] == 64
    assert asset["image_fit"] == "contain"
    assert Path(asset["image_path"]).read_bytes() == first.images[0].content
    assert provider.call_args.kwargs == {"quality": "medium", "aspect_ratio": "3:2"}
    assert "architectural photography" in provider.call_args.args[1]
    assert "No text" in provider.call_args.args[1]
    second = json.loads(
        tool.create_presentation_image("An operations desk", "executive").content
    )
    assert second["cached"] and second["data"]["image_path"] == asset["image_path"]
    assert provider.call_count == 1
    assert len(first.content) < 1500 and "base64" not in first.content


def test_provider_failure_never_claims_to_have_created_image(provider):
    provider.side_effect = ImageGenerationError("HTTP 402: account limit")
    result = PresentationImageTools("test-key").create_presentation_image("A desk")
    assert not json.loads(result.content)["ok"]
    assert not result.images
    assert "402" in result.content


def test_ppt_agent_receives_image_tool():
    agent = build_presentation_agent(
        user_id=None, session_id=None, message_id=None, debug_mode=False
    )
    assert any(
        "create_presentation_image" in getattr(tool, "functions", {})
        for tool in agent.tools
    )


def test_generated_image_really_embeds_in_deck(provider):
    asset = json.loads(
        PresentationImageTools("test-key")
        .create_presentation_image("An operations desk")
        .content
    )["data"]
    tool = PresentationTools(user_id=None, session_id=None, message_id=None)
    result = tool.create_presentation(
        "Pilot",
        [
            {
                "type": "title",
                "title": "A visual operations pilot",
                "image_path": asset["image_path"],
            },
            {
                "type": "image",
                "title": "Make the handoff visible",
                "image_path": asset["image_path"],
                "bullets": ["Keep the owner visible"],
                "caption": asset["caption"],
            },
        ],
    )
    assert result["ok"], result
    assert result["quality"]["browser_ok"]
    prs = Presentation(result["data"]["local_path"])
    for slide in prs.slides:
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        assert len(pictures) == 1
        assert abs(pictures[0].width / pictures[0].height - 1.5) < 0.002


@pytest.mark.parametrize("dimensions", [(300, 900), (900, 300), (400, 400)])
def test_contain_preserves_portrait_landscape_and_square(dimensions, tmp_path):
    path = tmp_path / "asset.jpg"
    Image.new("RGB", dimensions, "blue").save(path)
    tool = PresentationTools(user_id=None, session_id=None, message_id=None)
    result = tool.create_presentation(
        "Image fit",
        [
            {
                "type": "image",
                "title": "Keep the whole subject visible",
                "image_path": str(path),
                "image_fit": "contain",
                "bullets": ["Text remains beside the image"],
            }
        ],
    )
    assert result["ok"], result
    prs = Presentation(result["data"]["local_path"])
    picture = next(shape for shape in prs.slides[0].shapes if shape.shape_type == 13)
    assert abs(picture.width / picture.height - dimensions[0] / dimensions[1]) < 0.002
    assert (
        max(
            abs(picture.crop_left),
            abs(picture.crop_right),
            abs(picture.crop_top),
            abs(picture.crop_bottom),
        )
        <= 0.0001
    )


def test_webp_and_orientation_are_normalized(tmp_path):
    source = tmp_path / "image.webp"
    Image.new("RGB", (50, 100), "blue").save(source)
    slides = [{"image_path": str(source)}, {"image_path": str(source)}]
    prepare_slide_images(slides, tmp_path)
    assert slides[0]["image_path"].endswith(".png")
    assert slides[0]["image_path"] == slides[1]["image_path"]
    assert (slides[0]["image_width"], slides[0]["image_height"]) == (50, 100)


def test_invalid_fit_is_refused(tmp_path):
    tool = PresentationTools(user_id=None, session_id=None, message_id=None)
    result = tool.create_presentation(
        "Invalid fit",
        [
            {
                "type": "image",
                "title": "Bad fit",
                "image_path": "any.png",
                "image_fit": "stretch",
            }
        ],
    )
    assert not result["ok"]
    assert any(issue["field"] == "image_fit" for issue in result["issues"])


def test_exif_rotation_is_applied_before_image_dimensions(tmp_path):
    source = tmp_path / "rotated.jpg"
    image = Image.new("RGB", (60, 120), "blue")
    exif = image.getexif()
    exif[274] = 6
    image.save(source, exif=exif)
    slides = [{"image_path": str(source)}]
    prepare_slide_images(slides, tmp_path)
    assert (slides[0]["image_width"], slides[0]["image_height"]) == (120, 60)
    with Image.open(slides[0]["image_path"]) as prepared:
        assert prepared.getexif().get(274, 1) == 1
